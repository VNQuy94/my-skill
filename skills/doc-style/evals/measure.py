#!/usr/bin/env python3
"""Profile-independent layout metrics, for comparing runs with and without the skill.

Usage: measure.py <file.pdf|file.docx>   (run with the doc-style venv python)

Prints JSON: distinct text colors, distinct table header looks, fonts,
embedding, direct-formatting count, blank paragraphs, visual-structure counts.
"""
import json
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path


def rgb_hex(c):
    if c is None:
        return None
    c = tuple(float(x) for x in (c if isinstance(c, (list, tuple)) else (c,)) if isinstance(x, (int, float)))
    if len(c) == 1:
        c = (c[0],) * 3
    if len(c) == 4:
        k = c[3]
        c = tuple((1 - v) * (1 - k) for v in c[:3])
    if len(c) != 3:
        return None
    return "#" + "".join(f"{round(v * 255):02X}" for v in c)


def pdf_metrics(path: Path) -> dict:
    import pdfplumber

    colors, fills, fonts = Counter(), Counter(), Counter()
    pages = text_only_pages = 0
    replacement = 0
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            pages += 1
            for ch in page.chars:
                if ch["text"].strip():
                    colors[rgb_hex(ch.get("non_stroking_color"))] += 1
                    fonts[re.sub(r"^[A-Z]{6}\+", "", ch["fontname"])] += 1
                    replacement += ch["text"] == "�"
            for r in page.rects:
                if r.get("fill"):
                    fills[rgb_hex(r.get("non_stroking_color"))] += 1
            if not (page.rects or page.lines or page.images or page.curves):
                text_only_pages += 1
    emb = subprocess.run(["pdffonts", str(path)], capture_output=True, text=True).stdout.splitlines()[2:]
    not_embedded = [ln.split()[0] for ln in emb if len(ln.split()) >= 5 and ln.split()[-5] != "yes"]
    return {
        "pages": pages,
        "text_colors": len([c for c in colors if c]),
        "text_color_list": dict(colors.most_common()),
        "fill_colors": len(fills),
        "fill_color_list": dict(fills.most_common()),
        "font_faces": len(fonts),
        "fonts": dict(fonts.most_common()),
        "fonts_not_embedded": not_embedded,
        "missing_glyph_chars": replacement,
        "text_only_pages": text_only_pages,
    }


def docx_metrics(path: Path) -> dict:
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    direct = blank = bold_normal = 0
    styles = Counter()
    colors = Counter()
    for p in doc.paragraphs:
        styles[p.style.name if p.style is not None else "?"] += 1
        if not p.text.strip() and not p._p.xpath(".//w:drawing"):
            blank += 1
        runs = [r for r in p.runs if r.text.strip()]
        if runs and all(r.bold for r in runs) and (p.style is None or p.style.name == "Normal") and len(p.text.split()) <= 12:
            bold_normal += 1
        for r in p.runs:
            rpr = r._r.rPr
            if rpr is not None and any(rpr.find(qn(t)) is not None for t in ("w:color", "w:rFonts", "w:sz")):
                direct += 1
    for el in doc.element.iter(qn("w:color")):
        v = el.get(qn("w:val"))
        if v and v != "auto":
            colors[v.upper()] += 1
    header_fills = Counter()
    for t in doc.tables:
        if len(t.rows) >= 2 and len(t.columns) >= 2:
            fills = tuple(sorted({(s.get(qn("w:fill")) or "none").upper() for tc in t.rows[0]._tr.findall(qn("w:tc")) for s in tc.iter(qn("w:shd"))})) or ("none",)
            header_fills[fills] += 1
    lists = sum(1 for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None)
    images = len(doc.element.findall(".//" + qn("w:drawing")))
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "list_items": lists,
        "images": images,
        "heading_paragraphs": sum(v for k, v in styles.items() if k.startswith("Heading")),
        "fake_headings_bold_normal": bold_normal,
        "direct_formatted_runs": direct,
        "blank_paragraphs": blank,
        "inline_text_colors": len(colors),
        "inline_text_color_list": dict(colors),
        "distinct_table_header_looks": len(header_fills),
        "styles_used": dict(styles.most_common()),
    }


def main() -> None:
    path = Path(sys.argv[1])
    m = pdf_metrics(path) if path.suffix == ".pdf" else docx_metrics(path)
    print(json.dumps(m, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
