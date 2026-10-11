#!/usr/bin/env python3
"""Check a .docx against a doc-style profile.

Usage:
  lint_docx.py <file.docx> [--profile NAME|profile.json]

Errors (exit 1): direct color/font/size formatting, colors outside the
profile palette, unknown fonts, empty paragraphs used as spacing, tables whose
header rows look different from each other.
Warnings: walls of text, fake headings (bold Normal text), skipped heading
levels, too much bold.
If --profile is omitted, profile.json next to the file is used.
"""
import argparse
import re
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _env import ensure_venv, hex_to_rgb, load_profile, nearest_role, palette  # noqa: E402

ensure_venv("docx")
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
ALLOWED_EMPTY = {"BlockGap"}
HEADING_RE = re.compile(r"^Heading ?(\d)$")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--profile")
    a = ap.parse_args()
    path = Path(a.file)
    prof_arg = a.profile or (str(path.parent / "profile.json") if (path.parent / "profile.json").exists() else "modern")
    profile = load_profile(prof_arg)
    pal = palette(profile)
    fonts = {f for group in profile["fonts"].values() for f in group}

    doc = Document(str(path))
    errors: list[str] = []
    warns: list[str] = []

    parts = {"document": doc.element, "styles": doc.styles.element}
    try:
        parts["numbering"] = doc.part.numbering_part.element
    except Exception:
        pass

    # 1. Colors and fonts in the body, numbering, and the styles actually used
    used = {el.get(qn("w:val")) for tag in ("w:pStyle", "w:rStyle", "w:tblStyle") for el in doc.element.iter(qn(tag))} | {"Normal"}
    style_els = {s.get(qn("w:styleId")): s for s in doc.styles.element.findall(qn("w:style"))}
    stack = list(used)
    while stack:  # follow basedOn chains
        sid = stack.pop()
        el = style_els.get(sid)
        based = el.find(qn("w:basedOn")) if el is not None else None
        if based is not None and based.get(qn("w:val")) not in used:
            used.add(based.get(qn("w:val")))
            stack.append(based.get(qn("w:val")))
    doc_defaults = doc.styles.element.find(qn("w:docDefaults"))
    parts["styles"] = [style_els[s] for s in used if s in style_els] + ([doc_defaults] if doc_defaults is not None else [])

    roots = [(pname, r) for pname, rs in parts.items() for r in (rs if isinstance(rs, list) else [rs])]
    bad_colors: Counter = Counter()
    bad_fonts: set[str] = set()
    for pname, root in roots:
        for el in root.iter():
            tag = el.tag.replace(W, "")
            for attr in ("val", "fill", "color"):
                v = el.get(W + attr)
                if not v or v in ("auto", "nil"):
                    continue
                if (tag == "color" and attr == "val") or (tag == "shd" and attr == "fill") or (attr == "color" and tag in ("top", "bottom", "left", "right", "insideH", "insideV", "start", "end")):
                    if re.fullmatch(r"[0-9A-Fa-f]{6}", v) and nearest_role(hex_to_rgb(v), pal) is None:
                        bad_colors[f"#{v.upper()} ({pname})"] += 1
            if tag == "rFonts":
                for attr in ("ascii", "hAnsi", "cs"):
                    f = el.get(W + attr)
                    if f and f not in fonts:
                        bad_fonts.add(f"{f!r} ({pname})")
    errors += [f"font {f} is not in the profile" for f in sorted(bad_fonts)]
    for c, n in bad_colors.items():
        errors.append(f"color {c} is not a profile role color (used {n}×)")

    # 2. Paragraph-level checks
    body_run = 0
    max_body_run = 0
    prev_heading = 0
    bold_chars = total_chars = 0
    for i, para in enumerate(doc.paragraphs):
        style = para.style.style_id if para.style is not None else "Normal"
        text = para.text.strip()
        snippet = (text[:50] + "…") if len(text) > 50 else text

        if not text and not para._p.xpath(".//w:drawing") and style not in ALLOWED_EMPTY:
            errors.append(f"P{i}: empty paragraph (style {style}) — spacing must come from styles, not blank lines")

        for r in para.runs:
            n = len(r.text)
            total_chars += n
            if r.bold:
                bold_chars += n
            rpr = r._r.rPr
            if rpr is None:
                continue
            direct = [k for k, tag in (("color", "w:color"), ("font", "w:rFonts"), ("size", "w:sz"), ("highlight", "w:highlight"), ("shading", "w:shd")) if rpr.find(qn(tag)) is not None]
            if direct:
                errors.append(f"P{i}: direct {'/'.join(direct)} on run {r.text[:40]!r} — use a named style")

        ppr = para._p.pPr
        if ppr is not None and text:
            for tag in ("w:spacing", "w:ind", "w:jc"):
                if ppr.find(qn(tag)) is not None and ppr.find(qn("w:numPr")) is None:
                    warns.append(f"P{i}: direct paragraph {tag[2:]} on {snippet!r} — put it in the style")
                    break

        hm = HEADING_RE.match(para.style.name if para.style is not None else "")
        if hm:
            lvl = int(hm.group(1))
            if prev_heading and lvl > prev_heading + 1:
                warns.append(f"P{i}: heading level jumps from {prev_heading} to {lvl} at {snippet!r}")
            if lvl > profile["headings"]["max_level"]:
                warns.append(f"P{i}: heading level {lvl} exceeds max_level {profile['headings']['max_level']}")
            prev_heading = lvl
            body_run = 0
            continue

        if style == "Normal" and text:
            if para.runs and all(r.bold for r in para.runs if r.text.strip()) and len(text.split()) <= 12:
                warns.append(f"P{i}: bold Normal paragraph looks like a fake heading: {snippet!r}")
            words = len(text.split())
            if words > 150 and profile.get("lint", {}).get("wall_of_text", True):
                warns.append(f"P{i}: paragraph has {words} words — split it or use a list/table")
            body_run += 1
            max_body_run = max(max_body_run, body_run)
            if body_run == 6 and profile.get("lint", {}).get("wall_of_text", True):
                warns.append(f"P{i}: 6+ body paragraphs in a row with no heading/list/table/callout — consider a visual break")
        elif text:
            body_run = 0

    if total_chars and bold_chars / total_chars > 0.3:
        warns.append(f"bold text is {bold_chars / total_chars:.0%} of body text (keep under 30%)")

    # 3. Tables: every data table header must look the same
    header_looks = Counter()
    for t in doc.tables:
        rows = t.rows
        if len(rows) < 2 or len(t.columns) < 2:
            continue  # callouts and single-row stats
        cells = rows[0]._tr.findall(qn("w:tc"))
        fills = tuple(sorted({(shd.get(qn("w:fill")) or "none").upper() for tc in cells for shd in tc.iter(qn("w:shd"))})) or ("none",)
        header_looks[fills] += 1
    if len(header_looks) > 1:
        errors.append(f"data tables use {len(header_looks)} different header fills: {dict(header_looks)} — use t.table() for every table")

    for e in errors:
        print(f"ERROR  {e}")
    for w in warns:
        print(f"WARN   {w}")
    print(f"\n{path.name}: {len(errors)} error(s), {len(warns)} warning(s) · profile {profile['name']} · {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables")
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
